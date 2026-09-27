"""DOCX document parsing, batching, and text replacement engine."""

import asyncio
from dataclasses import dataclass
import io
import logging
import time
from typing import Any, Callable, Optional
import docx
import docx.text.paragraph

from app.core.llm_client import LLMClient
from app.core.run_aligner import count_text_diffs, update_paragraph_with_corrected_text
from app.core.track_changes import RevisionManager

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class DocumentElement:
    """Represents a processable text unit extracted from a DOCX file."""

    element_type: str  # "paragraph" or "table_cell"
    paragraph: docx.text.paragraph.Paragraph
    original_text: str

    @property
    def word_count(self) -> int:
        """Returns the word count of the element."""
        return len(self.original_text.split())

    def __iter__(self):
        yield self.element_type
        yield self.paragraph

    def __getitem__(self, idx: int):
        if idx == 0:
            return self.element_type
        if idx == 1:
            return self.paragraph
        raise IndexError("DocumentElement index out of range (0 or 1)")

    def __eq__(self, other: Any) -> bool:
        if isinstance(other, tuple) and len(other) == 2:
            return (self.element_type, self.paragraph) == other
        if isinstance(other, DocumentElement):
            return (
                self.element_type == other.element_type
                and self.paragraph == other.paragraph
                and self.original_text == other.original_text
            )
        return False


class DocxProcessor:
    """Handles parsing, batching, LLM correction, and in-place updating of DOCX documents."""

    def __init__(self, llm_client: LLMClient, concurrency_limit: int = 1) -> None:
        """Initializes the DocxProcessor.

        Args:
            llm_client: The LLM client used for text corrections.
            concurrency_limit: Maximum number of concurrent LLM requests (defaults to 1).
        """
        self.llm_client = llm_client
        self.concurrency_limit = concurrency_limit
        self.semaphore = asyncio.Semaphore(concurrency_limit)

    def extract_processable_paragraphs(self, doc: docx.Document) -> list[DocumentElement]:
        """Extracts all non-empty paragraphs from document body and tables.

        Table cells are deduplicated by underlying OpenXML element to avoid duplicate
        processing of merged cells.

        Args:
            doc: The python-docx Document object.

        Returns:
            List of DocumentElement instances to process.
        """
        elements: list[DocumentElement] = []

        # 1. Body paragraphs
        for p in doc.paragraphs:
            text = p.text
            if text and text.strip():
                elements.append(DocumentElement("paragraph", p, text))

        # 2. Table cells (deduplicating merged cells by XML node pointer _tc)
        visited_tcs = set()
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc in visited_tcs:
                        continue
                    visited_tcs.add(cell._tc)
                    for p in cell.paragraphs:
                        text = p.text
                        if text and text.strip():
                            elements.append(DocumentElement("table_cell", p, text))

        return elements

    def chunk_items_into_batches(
        self,
        elements: list[DocumentElement],
        max_batch_size: int = 6,
        max_batch_words: int = 250,
    ) -> list[list[DocumentElement]]:
        """Groups consecutive processable elements into batches by count and word limit.

        Args:
            elements: List of DocumentElement items to group.
            max_batch_size: Maximum number of items in a single batch.
            max_batch_words: Maximum aggregate words in a single batch.

        Returns:
            List of element batches.
        """
        batches: list[list[DocumentElement]] = []
        current_batch: list[DocumentElement] = []
        current_words = 0

        for elem in elements:
            words = elem.word_count
            if current_batch and (
                len(current_batch) >= max_batch_size
                or (current_words + words > max_batch_words and current_words > 0)
            ):
                batches.append(current_batch)
                current_batch = []
                current_words = 0

            current_batch.append(elem)
            current_words += words

        if current_batch:
            batches.append(current_batch)

        return batches

    async def process_document(
        self,
        docx_bytes: bytes,
        model_override: Optional[str] = None,
        progress_callback: Optional[Callable[[int, int, str, bool], None]] = None,
        cancel_check: Optional[Callable[[], bool]] = None,
        mode: str = "edit",
    ) -> tuple[bytes, dict[str, Any]]:
        """Processes a DOCX document in-place and returns updated bytes with statistics.

        Args:
            docx_bytes: Raw binary content of the source .docx document.
            model_override: Optional model name to override the default client model.
            progress_callback: Optional callback invoked with (processed, total, snippet, done).
            cancel_check: Optional callable returning True if processing was requested to stop.
            mode: Processing mode: "edit" (in-place replacement) or "suggest" (track changes).

        Returns:
            Tuple of (reconstructed_docx_bytes, metrics_dict).

        Raises:
            ValueError: If mode is not "edit" or "suggest".
            asyncio.CancelledError: If processing was cancelled.
            RuntimeError: If an unrecoverable LLM or processing error occurred.
        """
        if mode not in ("edit", "suggest"):
            raise ValueError(f"Invalid processing mode: '{mode}'. Must be 'edit' or 'suggest'.")

        start_time = time.time()
        input_stream = io.BytesIO(docx_bytes)
        doc = docx.Document(input_stream)

        revision_manager: Optional[RevisionManager] = None
        if mode == "suggest":
            revision_manager = RevisionManager(author="DocuMend")
            revision_manager.enable_track_revisions(doc)

        elements = self.extract_processable_paragraphs(doc)
        total_count = len(elements)
        body_count = sum(1 for e in elements if e.element_type == "paragraph")
        table_count = sum(1 for e in elements if e.element_type == "table_cell")
        total_words = sum(e.word_count for e in elements)

        logger.info(
            f"Extracted {total_count} processable elements ({total_words} words) from DOCX "
            f"[mode={mode}]: {body_count} body paragraph(s), {table_count} table cell(s)."
        )

        if total_count == 0:
            logger.warning("No processable text elements found in document.")
            empty_stats = {
                "elapsed_seconds": round(time.time() - start_time, 2),
                "total_words": 0,
                "words_per_second": 0.0,
                "tokens_per_second": None,
                "total_completion_tokens": None,
                "total_items": 0,
                "total_diffs": 0,
                "mode": mode,
            }
            if mode == "suggest":
                empty_stats["revisions_count"] = 0
            if progress_callback:
                progress_callback(0, 0, "No text found to process", True)
            return docx_bytes, empty_stats

        batches = self.chunk_items_into_batches(elements, max_batch_size=6, max_batch_words=250)
        logger.info(f"Grouped {total_count} elements into {len(batches)} batch(es) for processing.")

        processed_count = 0
        total_completion_tokens = 0
        total_gen_time = 0.0
        total_diffs = 0

        stop_processing_event = asyncio.Event()

        def is_cancelled() -> bool:
            return stop_processing_event.is_set() or bool(cancel_check and cancel_check())

        def _apply_correction(
            p: docx.text.paragraph.Paragraph,
            corr_text: str,
            orig_text: Optional[str] = None,
        ) -> int:
            if orig_text is None:
                orig_text = p.text
            diffs = count_text_diffs(orig_text, corr_text)
            if mode == "suggest" and revision_manager is not None:
                revision_manager.apply_revisions_to_paragraph(p, corr_text)
            else:
                update_paragraph_with_corrected_text(p, corr_text)
            return diffs

        async def _execute_single_paragraph(p: docx.text.paragraph.Paragraph) -> None:
            """Executes paragraph correction directly without acquiring semaphore.

            Assumes caller manages concurrency semaphore.
            """
            nonlocal processed_count, total_completion_tokens, total_gen_time, total_diffs
            if is_cancelled():
                return

            original_text = p.text
            snippet = (original_text[:60] + "...") if len(original_text) > 60 else original_text
            call_start = time.time()

            try:
                res = await self.llm_client.correct_text(
                    original_text, model_override=model_override
                )
                if is_cancelled():
                    return

                corrected_text = str(res)
                tokens = getattr(res, "completion_tokens", None)
                dur = getattr(res, "duration", time.time() - call_start)

                if tokens is not None:
                    total_completion_tokens += tokens
                total_gen_time += dur

                total_diffs += _apply_correction(p, corrected_text, original_text)
                processed_count += 1
                if progress_callback:
                    progress_callback(
                        processed_count,
                        total_count,
                        snippet,
                        processed_count >= total_count,
                    )
            except asyncio.CancelledError:
                raise
            except Exception as e:
                stop_processing_event.set()
                logger.error(f"Fatal error processing item #{processed_count + 1}: {e}")
                raise

        async def process_single_paragraph(p: docx.text.paragraph.Paragraph) -> None:
            """Wraps single paragraph execution in concurrency semaphore."""
            if is_cancelled():
                return
            async with self.semaphore:
                await _execute_single_paragraph(p)

        async def process_batch(batch: list[DocumentElement]) -> None:
            """Processes a batch of DocumentElement items."""
            nonlocal processed_count, total_completion_tokens, total_gen_time, total_diffs
            if is_cancelled():
                return

            if len(batch) == 1:
                await process_single_paragraph(batch[0].paragraph)
                return

            texts = [e.original_text for e in batch]
            batch_desc = f"batch of {len(batch)} items"

            async with self.semaphore:
                if is_cancelled():
                    return

                call_start = time.time()
                try:
                    corrected_texts, tokens, dur = await self.llm_client.correct_batch(
                        texts, model_override=model_override
                    )
                except asyncio.CancelledError:
                    raise
                except Exception as e:
                    stop_processing_event.set()
                    logger.error(f"Fatal server failure during {batch_desc}: {e}")
                    raise

                if tokens is not None:
                    total_completion_tokens += tokens
                total_gen_time += dur

                if is_cancelled():
                    return

                if corrected_texts is not None and len(corrected_texts) == len(batch):
                    for elem, corr in zip(batch, corrected_texts):
                        total_diffs += _apply_correction(elem.paragraph, corr, elem.original_text)
                        processed_count += 1
                        if progress_callback:
                            elem_snippet = (
                                (elem.original_text[:60] + "...")
                                if len(elem.original_text) > 60
                                else elem.original_text
                            )
                            progress_callback(
                                processed_count,
                                total_count,
                                elem_snippet,
                                processed_count >= total_count,
                            )
                else:
                    # Formatting mismatch fallback: process items directly while holding semaphore!
                    # BUG-01 FIX: Call _execute_single_paragraph directly to avoid semaphore deadlock.
                    logger.info(
                        f"Batched formatting mismatch; executing individual fallback for {batch_desc}"
                    )
                    for elem in batch:
                        if is_cancelled():
                            return
                        await _execute_single_paragraph(elem.paragraph)

        tasks = [asyncio.create_task(process_batch(b)) for b in batches]
        try:
            await asyncio.gather(*tasks)
        except Exception:
            stop_processing_event.set()
            for t in tasks:
                if not t.done():
                    t.cancel()
            raise

        if cancel_check and cancel_check():
            raise asyncio.CancelledError("Document processing was cancelled.")

        elapsed_total = time.time() - start_time
        words_per_sec = round(total_words / elapsed_total, 1) if elapsed_total > 0 else 0.0

        tokens_per_sec = None
        if total_completion_tokens > 0 and total_gen_time > 0:
            tokens_per_sec = round(total_completion_tokens / total_gen_time, 1)

        stats = {
            "elapsed_seconds": round(elapsed_total, 2),
            "total_words": total_words,
            "words_per_second": words_per_sec,
            "tokens_per_second": tokens_per_sec,
            "total_completion_tokens": (
                total_completion_tokens if total_completion_tokens > 0 else None
            ),
            "total_items": total_count,
            "total_diffs": total_diffs,
            "mode": mode,
        }
        if mode == "suggest" and revision_manager is not None:
            stats["revisions_count"] = revision_manager.revisions_count

        logger.info(
            f"Document processing completed: {total_words} words checked, {total_diffs} diff(s) in {elapsed_total:.2f}s "
            f"({words_per_sec} words/s"
            + (f", {tokens_per_sec} tok/s" if tokens_per_sec is not None else "")
            + ")."
        )

        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return output_stream.getvalue(), stats
