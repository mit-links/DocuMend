import docx
from app.core.run_aligner import (
    RunStyle,
    align_and_reconstruct_runs,
    update_paragraph_with_corrected_text,
)


def test_align_identical_text():
    runs_data = [
        ("Normal text ", RunStyle(bold=False)),
        ("bold text", RunStyle(bold=True)),
    ]
    result = align_and_reconstruct_runs(runs_data, "Normal text bold text")
    assert result == runs_data


def test_align_bold_correction():
    # Original: "Helo, this is a " (plain) + "test documnt" (bold)
    # Corrected: "Hello, this is a test document"
    runs_data = [
        ("Helo, this is a ", RunStyle(bold=None)),
        ("test documnt", RunStyle(bold=True)),
    ]
    corrected = "Hello, this is a test document"
    result = align_and_reconstruct_runs(runs_data, corrected)

    # Reconstructed text must match corrected
    reconstructed_text = "".join(text for text, _ in result)
    assert reconstructed_text == corrected

    # The bold section should cover "test document"
    bold_chunks = [text for text, style in result if style.bold is True]
    assert len(bold_chunks) > 0
    assert "document" in "".join(bold_chunks)


def test_update_paragraph_in_place():
    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Please read this ")
    r2 = p.add_run("importnt")
    r2.bold = True
    r2.italic = True
    r3 = p.add_run(" message.")

    modified = update_paragraph_with_corrected_text(p, "Please read this important message.")
    assert modified is True
    assert p.text == "Please read this important message."

    # Verify that bold and italic are preserved on the corrected word
    bold_italic_runs = [r for r in p.runs if r.bold and r.italic]
    assert len(bold_italic_runs) > 0
    assert "important" in "".join(r.text for r in bold_italic_runs)
