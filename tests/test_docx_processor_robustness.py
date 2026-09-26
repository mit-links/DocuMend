"""Robustness and edge-case unit tests for DocuMend's DocxProcessor.

Follows S/S/R naming convention and hermetic test principles.
Tests realistic document structures: empty paragraphs, tables, zero-text docs,
concurrency limits, and mid-flight cancellations.
"""

import asyncio
import io
import docx
import pytest

from app.core.docx_processor import DocxProcessor


class HermeticMockLLMClient:
    """Deterministic, hermetic mock LLM client for DocxProcessor tests."""

    def __init__(self, enable_batch: bool = True, delay: float = 0.0) -> None:
        self.enable_batch = enable_batch
        self.delay = delay
        self.current_concurrency = 0
        self.max_concurrency_observed = 0
        self.call_count = 0
        self.batch_call_count = 0

    async def correct_text(self, text: str, model_override=None):
        self.call_count += 1
        self.current_concurrency += 1
        self.max_concurrency_observed = max(
            self.max_concurrency_observed, self.current_concurrency
        )
        try:
            if self.delay > 0:
                await asyncio.sleep(self.delay)
            corrections = {
                "Helo world": "Hello world",
                "Speling error": "Spelling error",
                "Table hedr": "Table header",
                "Cell contnt": "Cell content",
            }
            return corrections.get(text.strip(), text)
        finally:
            self.current_concurrency -= 1

    async def correct_batch(self, texts, model_override=None):
        self.batch_call_count += 1
        self.current_concurrency += 1
        self.max_concurrency_observed = max(
            self.max_concurrency_observed, self.current_concurrency
        )
        try:
            if self.delay > 0:
                await asyncio.sleep(self.delay)
            if not self.enable_batch:
                return None, 10, 0.05
            corrections = {
                "Helo world": "Hello world",
                "Speling error": "Spelling error",
                "Table hedr": "Table header",
                "Cell contnt": "Cell content",
            }
            results = [corrections.get(t.strip(), t) for t in texts]
            return results, len(results) * 5, 0.05
        finally:
            self.current_concurrency -= 1


# ---------------------------------------------------------------------------
# Extraction & Batching Tests
# ---------------------------------------------------------------------------

def test_DocxProcessor_ExtractParagraphs_SkipsEmptyAndWhitespaceOnly():
    """Verify empty and whitespace paragraphs are skipped while retaining valid items."""
    doc = docx.Document()
    doc.add_paragraph("First valid paragraph.")
    doc.add_paragraph("")  # Empty
    doc.add_paragraph("   \n\t  ")  # Whitespace only
    doc.add_paragraph("Second valid paragraph.")

    # Table with an empty cell and a populated cell
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = ""
    table.rows[0].cells[1].paragraphs[0].text = "Valid cell text."

    processor = DocxProcessor(llm_client=HermeticMockLLMClient())
    items = processor.extract_processable_paragraphs(doc)

    assert len(items) == 3
    assert items[0].element_type == "paragraph"
    assert items[0].paragraph._p == doc.paragraphs[0]._p
    assert items[0].original_text == "First valid paragraph."
    assert items[1].element_type == "paragraph"
    assert items[1].paragraph._p == doc.paragraphs[3]._p
    assert items[1].original_text == "Second valid paragraph."
    assert items[2].element_type == "table_cell"
    assert items[2].paragraph.text == "Valid cell text."


def test_DocxProcessor_ChunkBatches_SplitsByCountAndWordLimits():
    """Verify batching respects max_batch_size and max_batch_words thresholds."""
    processor = DocxProcessor(llm_client=HermeticMockLLMClient())

    doc = docx.Document()
    for _ in range(5):
        doc.add_paragraph("word " * 10)
    items = processor.extract_processable_paragraphs(doc)

    # Test batch size constraint: max_batch_size=2
    batches_by_size = processor.chunk_items_into_batches(
        items, max_batch_size=2, max_batch_words=1000
    )
    assert len(batches_by_size) == 3
    assert len(batches_by_size[0]) == 2
    assert len(batches_by_size[1]) == 2
    assert len(batches_by_size[2]) == 1

    # Test word count constraint: max_batch_words=25 (each item has 10 words)
    batches_by_words = processor.chunk_items_into_batches(
        items, max_batch_size=10, max_batch_words=25
    )
    assert len(batches_by_words) == 3
    assert len(batches_by_words[0]) == 2  # 20 words
    assert len(batches_by_words[1]) == 2  # 20 words
    assert len(batches_by_words[2]) == 1  # 10 words


# ---------------------------------------------------------------------------
# Realistic Document Processing Tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_DocxProcessor_WhenDocumentIsEmpty_ReturnsOriginalBytesAndZeroStats():
    """Verify a document with zero text elements returns early with empty stats."""
    doc = docx.Document()
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    buffer = io.BytesIO()
    doc.save(buffer)
    input_bytes = buffer.getvalue()

    mock_client = HermeticMockLLMClient()
    processor = DocxProcessor(llm_client=mock_client)

    events = []
    output_bytes, stats = await processor.process_document(
        docx_bytes=input_bytes,
        progress_callback=lambda p, t, s, c: events.append((p, t, c)),
    )

    assert output_bytes == input_bytes
    assert stats["total_words"] == 0
    assert stats["total_items"] == 0
    assert mock_client.call_count == 0
    assert len(events) == 1
    assert events[0] == (0, 0, True)


@pytest.mark.asyncio
async def test_DocxProcessor_WithInterspersedEmptyParagraphs_PreservesDocumentStructure():
    """Verify document structure (indices and blank paragraphs) is preserved exactly."""
    doc = docx.Document()
    doc.add_paragraph("Helo world")
    doc.add_paragraph("")  # Empty spacer
    doc.add_paragraph("Speling error")
    doc.add_paragraph("")  # Trailing empty spacer

    buffer = io.BytesIO()
    doc.save(buffer)
    input_bytes = buffer.getvalue()

    mock_client = HermeticMockLLMClient()
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=2)

    output_bytes, stats = await processor.process_document(docx_bytes=input_bytes)
    output_doc = docx.Document(io.BytesIO(output_bytes))

    assert len(output_doc.paragraphs) == 4
    assert output_doc.paragraphs[0].text == "Hello world"
    assert output_doc.paragraphs[1].text == ""
    assert output_doc.paragraphs[2].text == "Spelling error"
    assert output_doc.paragraphs[3].text == ""
    assert stats["total_items"] == 2


@pytest.mark.asyncio
async def test_DocxProcessor_WithTables_CorrectsCellsAndPreservesDimensions():
    """Verify table cell contents are corrected while table structure is intact."""
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Table hedr"
    table.rows[0].cells[1].paragraphs[0].text = "Valid Header"
    table.rows[1].cells[0].paragraphs[0].text = ""  # Empty cell
    table.rows[1].cells[1].paragraphs[0].text = "Cell contnt"

    buffer = io.BytesIO()
    doc.save(buffer)

    mock_client = HermeticMockLLMClient()
    processor = DocxProcessor(llm_client=mock_client)

    output_bytes, stats = await processor.process_document(docx_bytes=buffer.getvalue())
    output_doc = docx.Document(io.BytesIO(output_bytes))

    out_table = output_doc.tables[0]
    assert len(out_table.rows) == 2
    assert len(out_table.columns) == 2
    assert out_table.rows[0].cells[0].paragraphs[0].text == "Table header"
    assert out_table.rows[0].cells[1].paragraphs[0].text == "Valid Header"
    assert out_table.rows[1].cells[0].paragraphs[0].text == ""
    assert out_table.rows[1].cells[1].paragraphs[0].text == "Cell content"
    assert stats["total_items"] == 3


@pytest.mark.asyncio
async def test_DocxProcessor_ConcurrencyLimit_StrictlyEnforcedBySemaphore():
    """Verify that parallel batch execution does not exceed concurrency_limit."""
    doc = docx.Document()
    for i in range(12):
        doc.add_paragraph(f"Helo world number {i}")

    buffer = io.BytesIO()
    doc.save(buffer)

    concurrency_limit = 2
    mock_client = HermeticMockLLMClient(delay=0.05)
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=concurrency_limit)

    output_bytes, stats = await processor.process_document(docx_bytes=buffer.getvalue())

    assert mock_client.max_concurrency_observed <= concurrency_limit
    assert stats["total_items"] == 12


@pytest.mark.asyncio
async def test_DocxProcessor_PartialBatchFailure_FallsBackToIndividualItems():
    """Verify that when batch parsing fails, processor falls back to single items."""
    doc = docx.Document()
    doc.add_paragraph("Helo world")
    doc.add_paragraph("Speling error")

    buffer = io.BytesIO()
    doc.save(buffer)

    # Batch disabled forces fallback path
    mock_client = HermeticMockLLMClient(enable_batch=False)
    processor = DocxProcessor(llm_client=mock_client)

    output_bytes, stats = await processor.process_document(docx_bytes=buffer.getvalue())
    output_doc = docx.Document(io.BytesIO(output_bytes))

    assert output_doc.paragraphs[0].text == "Hello world"
    assert output_doc.paragraphs[1].text == "Spelling error"
    assert mock_client.batch_call_count == 1
    assert mock_client.call_count == 2


@pytest.mark.asyncio
async def test_DocxProcessor_WhenCancelledMidProcessing_HaltsAndRaisesCancelledError():
    """Verify that cancelling while tasks are running terminates cleanly."""
    doc = docx.Document()
    for i in range(10):
        doc.add_paragraph(f"Paragraph {i}")

    buffer = io.BytesIO()
    doc.save(buffer)

    mock_client = HermeticMockLLMClient(delay=0.05)
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=1)

    is_cancelled = False

    def cancel_after_first_item(processed, total, snippet, is_complete):
        nonlocal is_cancelled
        if processed >= 1:
            is_cancelled = True

    with pytest.raises(asyncio.CancelledError):
        await processor.process_document(
            docx_bytes=buffer.getvalue(),
            progress_callback=cancel_after_first_item,
            cancel_check=lambda: is_cancelled,
        )
