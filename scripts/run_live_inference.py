"""Integration script testing DocuMend pipeline against the local LLM server."""
import asyncio
import io
import os
import sys

# Ensure repository root is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import docx
from app.core.docx_processor import DocxProcessor
from app.core.llm_client import LLMClient


async def main():
    print("Testing live inference with local LLM server at http://127.0.0.1:1234/v1...")
    client = LLMClient(base_url="http://127.0.0.1:1234/v1", api_key="not-needed")
    models = await client.list_models()
    print(f"Available models: {models}")
    selected_model = models[0]
    print(f"Using model: {selected_model}")

    with open("tests/sample_docs/test_doc.docx", "rb") as f:
        doc_bytes = f.read()

    orig_doc = docx.Document(io.BytesIO(doc_bytes))
    print("\n--- ORIGINAL DOCUMENT CONTENT ---")
    for i, p in enumerate(orig_doc.paragraphs):
        if p.text.strip():
            print(f"P{i+1}: {p.text}")
    for t_idx, table in enumerate(orig_doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip() for cell in row.cells]
            print(f"Table {t_idx+1} Row {r_idx+1}: {' | '.join(row_texts)}")

    processor = DocxProcessor(llm_client=client, concurrency_limit=1)

    def progress(processed, total, snippet, is_complete):
        print(f"Progress: [{processed}/{total}] {snippet} (complete={is_complete})")

    corrected_bytes, stats = await processor.process_document(
        docx_bytes=doc_bytes,
        model_override=selected_model,
        progress_callback=progress,
    )

    print(f"\n--- PERFORMANCE STATS ---")
    print(f"Elapsed Time: {stats['elapsed_seconds']}s")
    print(f"Words Checked: {stats['total_words']} ({stats['words_per_second']} words/sec)")
    if stats['tokens_per_second']:
        print(f"Generation Speed: {stats['tokens_per_second']} tokens/sec")

    with open("tests/sample_docs/test_doc_corrected.docx", "wb") as f:
        f.write(corrected_bytes)

    res_doc = docx.Document(io.BytesIO(corrected_bytes))
    print("\n--- CORRECTED DOCUMENT CONTENT ---")
    for i, p in enumerate(res_doc.paragraphs):
        if p.text.strip():
            print(f"P{i+1}: {p.text}")
            runs_info = [f"'{r.text}'(bold={r.bold}, italic={r.italic})" for r in p.runs]
            print(f"     Runs: {', '.join(runs_info)}")

    for t_idx, table in enumerate(res_doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip() for cell in row.cells]
            print(f"Table {t_idx+1} Row {r_idx+1}: {' | '.join(row_texts)}")


if __name__ == "__main__":
    asyncio.run(main())
