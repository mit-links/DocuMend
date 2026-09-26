import io
import pytest
import docx
from app.core.docx_processor import DocxProcessor
from scripts.create_test_doc import generate_sample_docx


class MockLLMClient:
    """Mock LLM client for deterministic unit testing."""

    def __init__(self, enable_batch: bool = True):
        self.enable_batch = enable_batch

    async def correct_text(self, text: str, model_override=None) -> str:
        corrections = {
            "Helo, this is a test documnt": "Hello, this is a test document",
            "I am writing many paragrafs, some of which have errors in speling.": "I am writing many paragraphs, some of which have errors in spelling.",
            "Einzelne Paragraphen sind in einer anderen Sprache, zum beispoiel deutsch.": "Einzelne Paragraphen sind in einer anderen Sprache, zum beispiel deutsch.",
            "Locaton": "Location",
        }
        return corrections.get(text, text)

    async def correct_batch(self, texts, model_override=None):
        if not self.enable_batch:
            return None, None, 0.0
        results = [await self.correct_text(t, model_override) for t in texts]
        return results, len(results) * 5, 0.05


@pytest.mark.asyncio
async def test_docx_processor_extraction_and_replacement(tmp_path):
    sample_file = tmp_path / "test.docx"
    generate_sample_docx(str(sample_file))

    with open(sample_file, "rb") as f:
        file_bytes = f.read()

    mock_client = MockLLMClient()
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=2)

    progress_events = []

    def on_progress(processed, total, snippet, is_complete):
        progress_events.append((processed, total, is_complete))

    output_bytes, stats = await processor.process_document(
        docx_bytes=file_bytes,
        progress_callback=on_progress,
    )

    assert len(output_bytes) > 0
    assert len(progress_events) > 0
    assert progress_events[-1][2] is True  # is_complete

    # Verify stats dictionary
    assert stats["total_words"] > 0
    assert stats["words_per_second"] >= 0
    assert stats["elapsed_seconds"] >= 0
    assert stats["total_items"] == len(progress_events)

    # Verify that the generated output docx can be parsed and has the corrected text
    output_doc = docx.Document(io.BytesIO(output_bytes))

    # Check body paragraphs
    para_texts = [p.text for p in output_doc.paragraphs if p.text.strip()]
    assert "Hello, this is a test document" in para_texts[0]
    assert "paragraphs" in para_texts[1]
    assert "spelling" in para_texts[1]
    assert "beispiel" in para_texts[2]  # German typo fixed, German preserved

    # Check table cells
    table = output_doc.tables[0]
    header_loc = table.rows[0].cells[1].paragraphs[0].text
    assert header_loc == "Location"  # "Locaton" corrected to "Location"


@pytest.mark.asyncio
async def test_docx_processor_cancellation(tmp_path):
    import asyncio
    sample_file = tmp_path / "test.docx"
    generate_sample_docx(str(sample_file))

    with open(sample_file, "rb") as f:
        file_bytes = f.read()

    mock_client = MockLLMClient()
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=2)

    # Cancel check returning True immediately
    with pytest.raises(asyncio.CancelledError):
        await processor.process_document(
            docx_bytes=file_bytes,
            cancel_check=lambda: True,
        )


@pytest.mark.asyncio
async def test_docx_processor_batch_fallback(tmp_path):
    sample_file = tmp_path / "test.docx"
    generate_sample_docx(str(sample_file))

    with open(sample_file, "rb") as f:
        file_bytes = f.read()

    # Mock client with batch disabled (forces fallback to individual items)
    mock_client = MockLLMClient(enable_batch=False)
    processor = DocxProcessor(llm_client=mock_client, concurrency_limit=2)

    output_bytes, stats = await processor.process_document(docx_bytes=file_bytes)
    assert len(output_bytes) > 0
    assert stats["total_items"] > 0

    output_doc = docx.Document(io.BytesIO(output_bytes))
    para_texts = [p.text for p in output_doc.paragraphs if p.text.strip()]
    assert "Hello, this is a test document" in para_texts[0]


